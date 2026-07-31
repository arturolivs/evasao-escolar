# -*- coding: utf-8 -*-
"""
Analisador de escrita da monografia — mede o que dá para medir.

Cinco blocos: ritmo e legibilidade, sintaxe PT-BR, clareza, conformidade ABNT
e marcadores de escrita automática. Não emite veredito: aponta trechos para
revisão humana, com o número do parágrafo para localizar no Word.

Uso (a partir da raiz do repositório):
    python .claude/skills/validar-escrita/analisar_escrita.py
    python .claude/skills/validar-escrita/analisar_escrita.py --doc outro.docx
    python .claude/skills/validar-escrita/analisar_escrita.py --bloco abnt
"""
from __future__ import annotations

import argparse
import io
import re
import statistics
import sys
import unicodedata
from collections import Counter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from docx import Document
from docx.shared import Cm

DOC_PADRAO = "documentos/TCC_Evasao_Escolar.docx"

# Parágrafos que não são texto corrido: não entram nas métricas de escrita.
ESTILOS_IGNORADOS = {"Caption", "TOC 1", "TOC 2", "TOC 3", "TOC 9", "Title"}
PREFIXOS_IGNORADOS = ("Figura ", "Quadro ", "Fonte:", "Tabela ")

LIMITE_FRASE_LONGA = 40      # palavras
LIMITE_FRASE_MUITO_LONGA = 55
LIMITE_PAR_LONGO = 8         # frases


# ---------------------------------------------------------------------------
# Carga e segmentação
# ---------------------------------------------------------------------------

def em_ingles(texto: str) -> bool:
    """Detecta o Abstract e os Keywords, que não entram na análise de PT-BR."""
    pal = [w.lower() for w in palavras(texto)]
    if not pal:
        return False
    marcas = {"the", "of", "and", "is", "with", "that", "which", "for", "from"}
    return sum(w in marcas for w in pal) / len(pal) > 0.12


def carregar(caminho: str):
    doc = Document(caminho)
    corpo = []          # (indice, texto) só de texto corrido em português
    todos = []          # (indice, estilo, texto)
    ref_inicio = None
    ingles = []
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        todos.append((i, p.style.name, txt))
        if re.fullmatch(r"REFERÊNCIAS\s*", txt, re.I):
            ref_inicio = i
        # A lista de referências e os apêndices não são prosa a avaliar:
        # ali «feature selection» é título de obra, não jargão do autor.
        if ref_inicio is not None and i > ref_inicio:
            continue
        if not txt or len(txt) < 40:
            continue
        if p.style.name in ESTILOS_IGNORADOS or p.style.name.startswith("Heading"):
            continue
        if txt.startswith(PREFIXOS_IGNORADOS):
            continue
        if em_ingles(txt):
            ingles.append(i)
            continue
        corpo.append((i, txt))
    if ingles:
        print(f"(¶{', ¶'.join(map(str, ingles))} em inglês — Abstract/Keywords "
              "fora da análise de PT-BR)")
    return doc, corpo, todos, ref_inicio


def frases(texto: str) -> list[str]:
    """Divide em frases protegendo abreviações e siglas com ponto."""
    t = re.sub(r"\b([A-Z])\.", r"\1<PT>", texto)          # iniciais
    t = t.replace("et al.", "et al<PT>")
    for abrev in ("p.", "n.", "v.", "ed.", "Prof.", "Dr.", "cf.", "etc."):
        t = t.replace(abrev, abrev[:-1] + "<PT>")
    partes = re.split(r"(?<=[.!?])\s+", t)
    return [p.replace("<PT>", ".").strip() for p in partes if p.strip()]


def palavras(texto: str) -> list[str]:
    return re.findall(r"\b[\wÀ-ÿ]+\b", texto)


def sem_acento(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s)
                   if unicodedata.category(c) != "Mn").lower()


# ---------------------------------------------------------------------------
# 1. Ritmo e legibilidade
# ---------------------------------------------------------------------------

def bloco_ritmo(corpo):
    print("=" * 78)
    print("1. RITMO E LEGIBILIDADE")
    print("=" * 78)

    todas, por_par = [], []
    for idx, txt in corpo:
        fs = frases(txt)
        tam = [len(palavras(f)) for f in fs]
        todas.extend((idx, f, n) for f, n in zip(fs, tam))
        por_par.append((idx, len(fs), sum(tam)))

    n_pal = [n for _, _, n in todas]
    if not n_pal:
        print("  sem texto corrido identificado.")
        return
    media = statistics.mean(n_pal)
    dp = statistics.pstdev(n_pal)
    print(f"  parágrafos de texto corrido : {len(corpo)}")
    print(f"  frases                      : {len(n_pal)}")
    print(f"  palavras/frase              : média {media:.1f} · dp {dp:.1f} "
          f"· mediana {statistics.median(n_pal):.0f} · máx {max(n_pal)}")
    print(f"  variação (dp/média)         : {dp/media:.2f}"
          "   [< 0,40 = ritmo uniforme; texto humano costuma variar mais]")

    longas = [(i, n, f) for i, f, n in todas if n >= LIMITE_FRASE_LONGA]
    print(f"\n  frases com ≥ {LIMITE_FRASE_LONGA} palavras: {len(longas)}"
          f"  ({100*len(longas)/len(n_pal):.1f}% do total)")
    for i, n, f in sorted(longas, key=lambda x: -x[1])[:8]:
        marca = "!!" if n >= LIMITE_FRASE_MUITO_LONGA else " ·"
        print(f"   {marca} ¶{i} ({n} palavras): {f[:110]}…")

    pars_longos = [(i, nf) for i, nf, _ in por_par if nf > LIMITE_PAR_LONGO]
    print(f"\n  parágrafos com > {LIMITE_PAR_LONGO} frases: {len(pars_longos)}")
    for i, nf in sorted(pars_longos, key=lambda x: -x[1])[:5]:
        print(f"    · ¶{i}: {nf} frases")


# ---------------------------------------------------------------------------
# 2. Sintaxe PT-BR
# ---------------------------------------------------------------------------

PARTICIPIO = r"[a-zà-ÿ]+(?:ado|ada|ados|adas|ido|ida|idos|idas)\b"

REGRAS_SINTAXE = [
    ("espaço duplo", r"[^\s] {2,}[^\s]"),
    ("espaço antes de pontuação", r"\s+[,;.!?](?:\s|$)"),
    ("falta espaço após pontuação", r"[,;:][A-Za-zÀ-ÿ]"),
    ("gerundismo (vai/vão estar + gerúndio)", r"\b(vai|vão|irá|irão)\s+estar\s+\w+ndo\b"),
    # «o mesmo» substituindo um substantivo (não o adjetivo «o mesmo critério»)
    ("«o mesmo» como pronome", r"\b(o|a|os|as)\s+mesmos?\s*(?:[,.;]|\s+(?:foi|foram|é|são|está|estão|deve|pode|apresenta))"),
    ("«através de» (prefira «por meio de»)", r"\batravés d[eo]s?\b"),
    ("«a nível de»", r"\ba n[íi]vel d[eo]\b"),
    ("«enquanto que»", r"\benquanto que\b"),
    ("primeira pessoa no corpo do trabalho",
     r"\b(eu|nós|nosso|nossa|nossos|nossas|meu|minha)\b"),
    ("coloquialismo", r"\b(daí|tipo assim|um monte|super|bastante bom|meio que)\b"),
    # «onde» como relativo de coisa não-lugar: o uso idiomático «mostra onde agir»
    # não entra, só a oração relativa «…, onde …»
    ("«onde» em oração relativa (conferir se é lugar)", r",\s*onde\b"),
]


def bloco_sintaxe(corpo, inicio_corpo: int | None = None):
    print("\n" + "=" * 78)
    print("2. SINTAXE E NORMA PT-BR")
    print("=" * 78)
    if inicio_corpo:
        print(f"  (pré-textuais até ¶{inicio_corpo} — agradecimentos e epígrafe usam")
        print("   primeira pessoa legitimamente e ficam fora desta contagem)\n")
    total = 0
    for nome, padrao in REGRAS_SINTAXE:
        achados = []
        for idx, txt in corpo:
            if inicio_corpo and idx < inicio_corpo and "primeira pessoa" in nome:
                continue
            for m in re.finditer(padrao, txt, re.I):
                ini = max(0, m.start() - 45)
                achados.append((idx, txt[ini:m.end() + 45].replace("\n", " ")))
        total += len(achados)
        status = "ok" if not achados else "XX"
        print(f"  {status} {nome}: {len(achados)}")
        for idx, ctx in achados[:3]:
            print(f"       ¶{idx}: …{ctx}…")
    # voz passiva (heurística)
    passivas = 0
    frases_tot = 0
    for _, txt in corpo:
        for f in frases(txt):
            frases_tot += 1
            if re.search(rf"\b(é|são|foi|foram|será|serão|sendo)\s+{PARTICIPIO}", f, re.I):
                passivas += 1
    pct = 100 * passivas / frases_tot if frases_tot else 0
    leitura = ("voz ativa predominante — bom para clareza" if pct < 15
               else "dentro do usual em texto acadêmico" if pct <= 30
               else "excesso de passiva: o texto esconde quem faz a ação")
    print(f"\n  voz passiva (heurística): {passivas} de {frases_tot} frases "
          f"({pct:.0f}%) — {leitura}")
    print(f"\n  total de ocorrências sinalizadas: {total}")


# ---------------------------------------------------------------------------
# 3. Clareza
# ---------------------------------------------------------------------------

JARGAO = ["feature", "features", "target", "baseline", "fold", "dataset",
          "overfitting", "pipeline", "ranking", "score", "outlier",
          "data leakage", "tuning", "boosting", "beeswarm", "waterfall"]

ABSTRATOS = ["no âmbito de", "no que tange", "no contexto de", "em termos de",
             "no sentido de", "de modo a", "com o intuito de", "de forma a",
             "faz-se necessário", "torna-se necessário", "cabe salientar"]


def bloco_clareza(corpo):
    print("\n" + "=" * 78)
    print("3. CLAREZA, SIMPLICIDADE E DIREÇÃO")
    print("=" * 78)

    texto = " ".join(t for _, t in corpo)
    pal = palavras(texto)
    n = len(pal) or 1

    nominal = [p for p in pal if re.search(r"(ção|ções|mento|mentos|dade|dades)$", p, re.I)]
    print(f"  palavras de texto corrido      : {n}")
    print(f"  nominalizações (-ção/-mento/…) : {len(nominal)} "
          f"({100*len(nominal)/n:.1f}%)   [acima de ~9% o texto fica abstrato]")

    print("\n  jargão técnico em inglês (o orientador pediu linguagem de gestor):")
    print("    Uso legítimo: glosa entre parênteses após o termo em português, e")
    print("    identificador de código (src.features). O resto é candidato a troca.")
    revisar, legitimos = [], 0
    for j in JARGAO:
        for i, t in corpo:
            for m in re.finditer(rf"\b{j}\b", t, re.I):
                antes = t[max(0, m.start() - 30):m.start()]
                if "(" in antes and ")" in t[m.end():m.end() + 3]:
                    legitimos += 1          # glosa: «gradiente impulsionado (gradient boosting)»
                elif re.search(r"src[\._]$|`$", antes):
                    legitimos += 1          # identificador de código
                else:
                    ctx = t[max(0, m.start() - 60):m.end() + 60].replace("\n", " ")
                    revisar.append((i, j, ctx))
    print(f"    {legitimos} ocorrência(s) em uso legítimo")
    print(f"    {'ok' if not revisar else 'XX'} {len(revisar)} candidato(s) a troca:")
    for i, j, ctx in revisar[:6]:
        print(f"       ¶{i} «{j}»: …{ctx}…")

    print("\n  perífrases que alongam sem informar:")
    achou_abs = False
    for a in ABSTRATOS:
        hits = [i for i, t in corpo if a in t.lower()]
        if hits:
            achou_abs = True
            print(f"    XX «{a}»: {len(hits)} — ¶{', ¶'.join(str(i) for i in hits[:5])}")
    if not achou_abs:
        print("    ok nenhuma da lista")

    # encadeamento de "que"
    print("\n  encadeamento de «que» (3+ na mesma frase):")
    pesadas = []
    for i, t in corpo:
        for f in frases(t):
            c = len(re.findall(r"\bque\b", f, re.I))
            if c >= 3:
                pesadas.append((i, c, f))
    print(f"    {len(pesadas)} frase(s)")
    for i, c, f in sorted(pesadas, key=lambda x: -x[1])[:4]:
        print(f"     · ¶{i} ({c}×): {f[:105]}…")

    # repetição de palavra rara na mesma frase
    print("\n  palavra significativa repetida na mesma frase:")
    stop = set("de da do das dos a o as os e ou que em no na nos nas um uma para "
               "com por se ao aos à às como mais não é são foi ser entre sobre "
               "isso este esta esse essa seu sua".split())
    reps = []
    for i, t in corpo:
        for f in frases(t):
            c = Counter(w.lower() for w in palavras(f)
                        if len(w) > 5 and w.lower() not in stop)
            for w, k in c.items():
                if k >= 3:
                    reps.append((i, w, k))
    print(f"    {len(reps)} caso(s)")
    for i, w, k in reps[:5]:
        print(f"     · ¶{i}: «{w}» {k}×")


# ---------------------------------------------------------------------------
# 4. ABNT
# ---------------------------------------------------------------------------

def bloco_abnt(doc, corpo, todos, ref_inicio):
    print("\n" + "=" * 78)
    print("4. CONFORMIDADE ABNT")
    print("=" * 78)

    # -- formatação da página: o MODELO DA PUC-SP prevalece sobre a leitura
    # genérica da NBR 14724 (3/3/2/2). O modelo usa 3,0/2,5/3,0/2,5 e é ele que
    # a banca cobra — por isso a comparação é contra o arquivo do modelo.
    def cm(v):
        return v / Cm(1) if v is not None else None

    s = doc.sections[0]
    atual = {n: cm(v) for n, v in (("esquerda", s.left_margin), ("superior", s.top_margin),
                                   ("direita", s.right_margin), ("inferior", s.bottom_margin))}
    modelo = None
    for cand in ("Modelo TCC Engenharia de Software 2022.docx",
                 "../Modelo TCC Engenharia de Software 2022.docx"):
        try:
            m = Document(cand).sections[0]
            modelo = {"esquerda": cm(m.left_margin), "superior": cm(m.top_margin),
                      "direita": cm(m.right_margin), "inferior": cm(m.bottom_margin)}
            print(f"  margens — comparadas ao modelo institucional ({cand}):")
            break
        except Exception:
            continue
    if modelo is None:
        modelo = {"esquerda": 3.0, "superior": 3.0, "direita": 2.0, "inferior": 2.0}
        print("  margens — modelo institucional não encontrado; usando NBR 14724 genérica:")
    for nome, valor in atual.items():
        alvo = modelo[nome]
        ok = valor is not None and alvo is not None and abs(valor - alvo) < 0.2
        print(f"    {'ok' if ok else 'XX'} {nome}: {valor:.1f} cm (modelo {alvo:.1f} cm)")

    # -- corpo do texto
    try:
        est = doc.styles["normal1"]
        tam = est.font.size.pt if est.font.size else None
        esp = est.paragraph_format.line_spacing
        print(f"\n  estilo do corpo «normal1»: fonte {tam} pt · entrelinha {esp} "
              "(esperado 12 pt e 1,5)")
    except KeyError:
        print("\n  ?? estilo «normal1» não encontrado — conferir manualmente")

    # -- citações e referências
    texto = "\n".join(t for _, _, t in todos[:ref_inicio] if t) if ref_inicio else \
            "\n".join(t for _, t in corpo)

    # Sobrenomes citados. Duas formas ABNT, com armadilhas reais:
    #   (CHAPMAN et al., 2000)  → «et al.» em minúscula quebra classe só-maiúscula
    #   (McKINNEY, 2010)        → grafia mista
    #   Ribeiro, Singh e Guestrin (2016) → co-autores em Title Case
    citadas: set[str] = set()
    for m in re.finditer(r"\(([^)]{2,140}?),\s*(?:19|20)\d{2}[a-c]?", texto):
        for parte in m.group(1).split(";"):
            tok = re.match(r"\s*((?:Mc|Mac|D[aeo]s?|Van|Von)?\s*[A-ZÀ-Ú][A-Za-zÀ-ÿ\-']+)",
                           parte.strip())
            if tok:
                citadas.add(sem_acento(tok.group(1).strip()))
    for m in re.finditer(
        r"\b((?:[A-ZÀ-Ú][a-zà-ÿ]+)(?:(?:,|\s+e)\s+[A-ZÀ-Ú][a-zà-ÿ]+)*)"
        r"(?:\s+et al\.)?\s*\((?:19|20)\d{2}[a-c]?\)", texto):
        for nome in re.split(r",|\s+e\s+", m.group(1)):
            if len(nome.strip()) > 2:
                citadas.add(sem_acento(nome.strip()))
    citadas = {c for c in citadas if len(c) > 2}

    referencias, ordem = [], []
    if ref_inicio is not None:
        for i, estilo, txt in todos[ref_inicio + 1:]:
            if not txt or len(txt) < 25:
                continue
            if estilo.startswith("Heading") or re.match(r"^(AP[ÊE]NDICE|ANEXO)", txt, re.I):
                break
            referencias.append((i, txt))
            m = re.match(r"^((?:Mc|Mac)?[A-ZÀ-Ú][A-Za-zÀ-ÿ\-']+)", txt)
            ordem.append(sem_acento(m.group(1)) if m else sem_acento(txt[:18]))

    print(f"\n  entradas na lista de referências: {len(referencias)}")
    print(f"  sobrenomes citados no texto: {len(citadas)}")

    # Uma citação está coberta se o sobrenome aparece em QUALQUER entrada —
    # co-autor não abre entrada própria, mas consta no corpo da referência.
    blob = sem_acento(" | ".join(t for _, t in referencias))
    orfas = sorted(c for c in citadas if c not in blob)
    print(f"\n  {'ok' if not orfas else 'XX'} citações sem entrada nas referências: "
          f"{len(orfas)}")
    for c in orfas[:8]:
        print(f"       · {c.upper()}")

    texto_citacoes = sem_acento(texto)
    nao_citadas = sorted({r for r in ordem if r not in texto_citacoes})
    print(f"  {'ok' if not nao_citadas else 'XX'} referências não citadas no texto: "
          f"{len(nao_citadas)}")
    for r in nao_citadas[:8]:
        print(f"       · {r.upper()}")

    fora = [(a, b) for a, b in zip(ordem, ordem[1:]) if a > b]
    print(f"\n  {'ok' if not fora else 'XX'} ordem alfabética das referências: "
          f"{len(fora)} quebra(s)")
    for a, b in fora[:5]:
        print(f"       · «{a.upper()}» antes de «{b.upper()}»")

    sem_ponto = [i for i, t in referencias if not t.rstrip().endswith(".")]
    print(f"  {'ok' if not sem_ponto else 'XX'} referências sem ponto final: "
          f"{len(sem_ponto)}"
          + (f" — ¶{', ¶'.join(map(str, sem_ponto[:5]))}" if sem_ponto else ""))

    # -- citação direta longa (NBR 10520: > 3 linhas → recuo 4 cm, fonte menor, simples)
    # Recuo de 4 cm com texto substancial. Faixa 3,5–5,5 cm exclui os blocos de
    # 8 cm da folha de rosto, que não são citação.
    print("\n  citações diretas longas (NBR 10520: recuo 4 cm, fonte menor, simples):")
    longas = []
    for p in doc.paragraphs:
        t = p.text.strip()
        rec = p.paragraph_format.left_indent
        if rec is not None and Cm(3.5) <= rec <= Cm(5.5) and len(t) > 150:
            longas.append((p, t, rec))
    print(f"    blocos de citação recuada: {len(longas)}")
    for p, t, rec in longas[:5]:
        tam = p.style.font.size.pt if p.style.font.size else None
        esp = p.paragraph_format.line_spacing or (
            p.style.paragraph_format.line_spacing)
        alerta = []
        if tam and tam >= 12:
            alerta.append("fonte não reduzida")
        if esp and esp > 1.05:
            alerta.append(f"entrelinha {esp} (esperado 1,0)")
        marca = "XX" if alerta else "ok"
        print(f"     {marca} recuo {rec/Cm(1):.1f} cm · fonte {tam} pt"
              + (f" — {'; '.join(alerta)}" if alerta else ""))
        print(f"        {t[:80]}…")
    embutidas = [(i, t) for i, t in corpo
                 if re.search(r"[“\"][^”\"]{280,}[”\"]", t)]
    print(f"    {'ok' if not embutidas else 'XX'} citações longas ainda no corpo: "
          f"{len(embutidas)}")
    for i, _ in embutidas[:3]:
        print(f"       · ¶{i} — mover para bloco recuado")

    # -- legenda acima / fonte abaixo
    print("\n  legenda de figura acima e «Fonte:» abaixo:")
    seq = [(i, t) for i, _, t in todos if t.startswith(("Figura ", "Quadro ", "Fonte:"))]
    problemas = 0
    for (i1, t1), (i2, t2) in zip(seq, seq[1:]):
        if t1.startswith(("Figura ", "Quadro ")) and not t2.startswith("Fonte:"):
            if not t2.startswith(("Figura ", "Quadro ")):
                problemas += 1
    print(f"    {'ok' if problemas == 0 else '??'} pares legenda/fonte fora do padrão: "
          f"{problemas}   [conferir no Word: o par pode estar separado pela imagem]")


# ---------------------------------------------------------------------------
# 5. Marcadores de escrita automática
# ---------------------------------------------------------------------------

CONECTORES = ["além disso", "portanto", "dessa forma", "desse modo", "nesse sentido",
              "por outro lado", "em suma", "em resumo", "vale destacar",
              "vale ressaltar", "é importante ressaltar", "cabe destacar",
              "por fim", "assim", "ademais", "outrossim", "com efeito"]

SUPERLATIVOS = ["fundamental", "crucial", "robusto", "robusta", "significativo",
                "significativa", "amplamente", "notavelmente", "extremamente",
                "essencial", "poderoso", "inovador", "abrangente", "holístico",
                "aprofundado", "valioso", "relevante"]


def bloco_ia(corpo):
    print("\n" + "=" * 78)
    print("5. MARCADORES DE ESCRITA AUTOMÁTICA (indícios, não veredito)")
    print("=" * 78)
    print("  Nenhum item abaixo prova autoria de IA. São padrões que modelos de")
    print("  linguagem produzem com frequência acima da média humana — e que também")
    print("  aparecem em texto acadêmico formal legítimo. Use como lista de revisão.\n")

    texto = " ".join(t for _, t in corpo).lower()
    n_pal = len(palavras(texto)) or 1
    todas_frases = [(i, f) for i, t in corpo for f in frases(t)]

    # a) uniformidade de ritmo
    tam = [len(palavras(f)) for _, f in todas_frases]
    dp_rel = statistics.pstdev(tam) / statistics.mean(tam) if tam else 0
    sinal = "indício" if dp_rel < 0.40 else "ok"
    print(f"  a) {sinal:8} variação do tamanho das frases: {dp_rel:.2f}")
    print("               [< 0,40 = ritmo homogêneo, marca típica de geração]")

    # b) conectores de escada
    print("\n  b) conectores formulaicos:")
    tot_con = 0
    for c in CONECTORES:
        k = len(re.findall(rf"\b{c}\b", texto))
        tot_con += k
        if k >= 3:
            print(f"       «{c}»: {k}×")
    dens = 1000 * tot_con / n_pal
    print(f"     total {tot_con} em {n_pal} palavras ({dens:.1f} por mil)"
          f"   [{'indício' if dens > 6 else 'ok'}: acima de 6/mil soa mecânico]")

    # c) superlativos vazios
    print("\n  c) adjetivos de reforço sem medida ao lado:")
    tot_sup = 0
    for s in SUPERLATIVOS:
        k = len(re.findall(rf"\b{s}\b", texto))
        tot_sup += k
        if k >= 3:
            print(f"       «{s}»: {k}×")
    d_sup = 1000 * tot_sup / n_pal
    print(f"     total {tot_sup} ({d_sup:.1f} por mil)"
          f"   [{'indício' if d_sup > 4 else 'ok'}]")

    # d) tricolon
    tri = sum(1 for _, f in todas_frases
              if re.search(r"\b\w+,\s+\w+\s+e\s+\w+\b", f))
    pct_tri = 100 * tri / len(todas_frases) if todas_frases else 0
    print(f"\n  d) {'indício' if pct_tri > 12 else 'ok':8} listas de três termos "
          f"(«X, Y e Z»): {tri} frases ({pct_tri:.1f}%)")

    # e) "não apenas ... mas também"
    par = len(re.findall(r"não (apenas|só)[^.]{0,80}mas também", texto))
    print(f"  e) {'indício' if par > 2 else 'ok':8} «não apenas… mas também»: {par}×")

    # f) travessão
    trav = texto.count("—")
    d_trav = 1000 * trav / n_pal
    print(f"  f) {'indício' if d_trav > 4 else 'ok':8} travessões «—»: {trav} "
          f"({d_trav:.1f} por mil)")

    # g) aberturas repetidas
    print("\n  g) aberturas de frase repetidas (3+ primeiras palavras iguais):")
    ab = Counter(" ".join(palavras(f)[:3]).lower() for _, f in todas_frases
                 if len(palavras(f)) >= 4)
    rep = [(a, k) for a, k in ab.most_common(8) if k >= 4]
    if rep:
        for a, k in rep:
            print(f"       «{a}…»: {k}×")
    else:
        print("       ok nenhuma abertura repetida 4+ vezes")

    # h) parágrafos de tamanho uniforme
    tam_par = [len(palavras(t)) for _, t in corpo]
    dp_par = statistics.pstdev(tam_par) / statistics.mean(tam_par) if tam_par else 0
    print(f"\n  h) {'indício' if dp_par < 0.35 else 'ok':8} variação do tamanho dos "
          f"parágrafos: {dp_par:.2f}")

    print("\n  --- leitura do bloco ---")
    print("  Vários «indício» juntos justificam uma passagem de reescrita; um ou dois,")
    print("  não. O que não se mede aqui e vale mais que tudo isto: se o texto conta o")
    print("  que ACONTECEU neste projeto (decisões, becos sem saída, números próprios),")
    print("  ele é seu, independentemente de quem redigiu a frase.")


# ---------------------------------------------------------------------------

BLOCOS = {"ritmo": 1, "sintaxe": 2, "clareza": 3, "abnt": 4, "ia": 5}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--doc", default=DOC_PADRAO)
    ap.add_argument("--bloco", choices=sorted(BLOCOS), default=None,
                    help="roda só um bloco (padrão: todos)")
    args = ap.parse_args()

    doc, corpo, todos, ref_inicio = carregar(args.doc)
    inicio_corpo = next((i for i, est, _ in todos if est == "Heading 1"), None)
    print(f"\ndocumento: {args.doc}")
    print(f"parágrafos: {len(todos)} · texto corrido: {len(corpo)} · "
          f"corpo começa em ¶{inicio_corpo} · referências a partir de ¶{ref_inicio}\n")

    quais = [args.bloco] if args.bloco else list(BLOCOS)
    if "ritmo" in quais:
        bloco_ritmo(corpo)
    if "sintaxe" in quais:
        bloco_sintaxe(corpo, inicio_corpo)
    if "clareza" in quais:
        bloco_clareza(corpo)
    if "abnt" in quais:
        bloco_abnt(doc, corpo, todos, ref_inicio)
    if "ia" in quais:
        bloco_ia(corpo)
    print()


if __name__ == "__main__":
    main()
