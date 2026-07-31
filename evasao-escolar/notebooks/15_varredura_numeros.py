# -*- coding: utf-8 -*-
"""Varredura: cada numero do TCC contra o artefato que o gera.

Nao le o valor "de memoria": recalcula a partir dos parquets/CSVs do projeto e
compara com o que esta escrito no .docx.
"""
import sys, io, os, re
import numpy as np
import pandas as pd
from docx import Document

sys.path.insert(0, r"C:\Users\Suporte\Documents\monografia\evasao-escolar")
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

RAIZ = r"C:\Users\Suporte\Documents\monografia\evasao-escolar"
DOC = r"C:\Users\Suporte\Documents\monografia\documentos\TCC_Evasao_Escolar.docx"

doc = Document(DOC)
PARS = [p.text for p in doc.paragraphs]
TXT = "\n".join(PARS)
TABS = doc.tables

res = {"ok": 0, "div": [], "na": []}


def checa(rotulo, no_texto, calculado, tol=0.011, fonte=""):
    """Compara valor do texto x recalculado. tol relativa a magnitude."""
    if calculado is None:
        res["na"].append((rotulo, no_texto, fonte))
        print(f"  ?? {rotulo}: texto={no_texto} | sem fonte reproduzivel")
        return
    if isinstance(no_texto, str):
        no_texto = float(no_texto.replace(".", "").replace(",", "."))
    dif = abs(no_texto - calculado)
    limite = max(tol, abs(calculado) * tol)
    if dif <= limite:
        res["ok"] += 1
        print(f"  ok {rotulo}: {no_texto} ≈ {calculado:.4g}")
    else:
        res["div"].append((rotulo, no_texto, calculado, fonte))
        print(f"  XX {rotulo}: TEXTO={no_texto}  CALCULADO={calculado:.4g}   [{fonte}]")


def no_texto(padrao, grupo=1):
    m = re.search(padrao, TXT)
    return m.group(grupo) if m else None


# =========================================================== FONTES
painel = pd.read_parquet(os.path.join(RAIZ, "data/interim/painel_escola_ano_pe_estadual_em.parquet"))
feat = pd.read_parquet(os.path.join(RAIZ, "data/processed/features.parquet"))
taxas = pd.read_parquet(os.path.join(RAIZ, "data/interim/taxas_rendimento_pe_estadual_em.parquet"))
est = pd.read_csv(os.path.join(RAIZ, "reports/estatisticas_features.csv"))
meso = pd.read_csv(os.path.join(RAIZ, "reports/estatisticas_mesorregiao.csv"))
cvrep = pd.read_csv(os.path.join(RAIZ, "reports/metricas_cv_repetida.csv"))
mxgb = pd.read_csv(os.path.join(RAIZ, "reports/metricas_xgboost.csv"))
ss13 = pd.read_csv(os.path.join(RAIZ, "reports/metricas_ss13_cenarios.csv"))
ralunos = pd.read_csv(os.path.join(RAIZ, "reports/residuos_transicoes_alunos.csv"))
rgrupo = pd.read_csv(os.path.join(RAIZ, "reports/residuos_grupo_temporal.csv"))
rdiag = pd.read_csv(os.path.join(RAIZ, "reports/residuos_top20_temporal.csv"))
shap_imp = pd.read_csv(os.path.join(RAIZ, "reports/shap_importancia.csv"))

print("=" * 78)
print("1. BASE HISTORICA E CONJUNTO DE MODELAGEM")
print("=" * 78)
checa("painel: registros", 2392, len(painel), fonte="painel parquet")
checa("painel: escolas", 806, painel["CO_ENTIDADE"].nunique(), fonte="painel parquet")
n3 = (painel.groupby("CO_ENTIDADE")["NU_ANO_CENSO"].nunique() == 3).mean() * 100
checa("painel: % em 3 anos", 97.4, n3, tol=0.05, fonte="painel parquet")
checa("painel: municipios", 185, painel["NO_MUNICIPIO"].nunique(), fonte="painel parquet")
urb = (painel["TP_LOCALIZACAO"] == 1).mean() * 100 if "TP_LOCALIZACAO" in painel else None
checa("painel: % urbanas", 86, urb, tol=0.6, fonte="painel parquet")
med_mat = painel["QT_MAT_MED"].median() if "QT_MAT_MED" in painel else None
checa("painel: mediana matriculas (~350)", 350, med_mat, tol=0.10, fonte="painel parquet")

checa("modelagem: observacoes", 1586, len(feat), fonte="features parquet")
checa("modelagem: escolas", 801, feat["CO_ENTIDADE"].nunique(), fonte="features parquet")
n2024 = (painel["NU_ANO_CENSO"] == 2024).sum()
checa("Quadro 5: linhas de 2024 descartadas como origem", 791, n2024, fonte="painel parquet")
checa("Quadro 5: pares potenciais", 1601, len(painel) - n2024, fonte="painel parquet")
checa("Quadro 5: pares perdidos", 15, (len(painel) - n2024) - len(feat), fonte="painel parquet")

print()
print("=" * 78)
print("2. O PROBLEMA (taxas da rede, concentracao)")
print("=" * 78)
tx = taxas.groupby("NU_ANO_CENSO")["TAXA_ABND_MED"].mean()
for ano, esperado in [(2022, 2.05), (2023, 1.37), (2024, 0.96)]:
    checa(f"taxa media da rede {ano}", esperado, float(tx.get(ano, np.nan)), tol=0.02,
          fonte="taxas parquet")

# concentracao em 2024: alunos que abandonaram = taxa/100 * matricula
p24 = painel[painel["NU_ANO_CENSO"] == 2024][["CO_ENTIDADE", "QT_MAT_MED"]]
t24 = taxas[taxas["NU_ANO_CENSO"] == 2024][["CO_ENTIDADE", "TAXA_ABND_MED"]]
j = p24.merge(t24, on="CO_ENTIDADE", how="inner")
j["alunos"] = j["TAXA_ABND_MED"] / 100 * j["QT_MAT_MED"]
total = j["alunos"].sum()
checa("2024: alunos que abandonaram", 2043, total, tol=0.01, fonte="painel x taxas")
checa("2024: % escolas com abandono zero", 67, (j["TAXA_ABND_MED"] == 0).mean() * 100,
      tol=0.02, fonte="painel x taxas")
k = int(round(len(j) * 0.10))
checa("2024: nº de escolas no decil superior", 79, k, fonte="painel x taxas")
top = j.nlargest(k, "alunos")["alunos"].sum()
checa("2024: % dos evadidos no decil superior", 76, top / total * 100, tol=0.02,
      fonte="painel x taxas")

# Dois recortes distintos, ambos corretos — a confusao entre eles ja gerou uma
# duvida registrada como pendencia (PLANO_DE_FINALIZACAO, item 5.2):
#   2.043 = rede estadual completa em 2024 (791 escolas com taxa e matricula)
#   2.020 = coorte de modelagem 2023->2024 (789 escolas que formam par t->t+1)
# A diferenca sao as 2 escolas sem par no ano seguinte, que somam 23 alunos.
coorte24 = ralunos[ralunos["transicao"] == "2023-2024"]
checa("2024: alunos na coorte de modelagem", 2020, coorte24["alunos_real"].sum(),
      tol=0.01, fonte="residuos_transicoes_alunos")
checa("2024: escolas na coorte de modelagem", 789, len(coorte24),
      fonte="residuos_transicoes_alunos")
checa("2024: alunos fora da coorte (rede - coorte)", 23,
      total - coorte24["alunos_real"].sum(), tol=0.05, fonte="painel x taxas")

print()
print("=" * 78)
print("3. TEMPO INTEGRAL E INFRAESTRUTURA (Cap. 1 e 3)")
print("=" * 78)
if "QT_MAT_MED_INT" in painel.columns:
    pi = painel.assign(pct=lambda d: np.where(d["QT_MAT_MED"] > 0,
                                              d["QT_MAT_MED_INT"] / d["QT_MAT_MED"] * 100, np.nan))
    for ano, esperado in [(2022, 56.7), (2024, 64.9)]:
        v = pi[pi["NU_ANO_CENSO"] == ano]["pct"].mean()
        checa(f"% matriculas em tempo integral {ano}", esperado, v, tol=0.02, fonte="painel parquet")
else:
    checa("% escolas com tempo integral", 56.7, None, fonte="coluna ausente no painel")
# definicao do nb02: media dos 13 itens de infra, ANO 2024, urbana x rural
# por TP_LOCALIZACAO (sem separar 'diferenciada')
import re as _re
_src = open(os.path.join(RAIZ, "notebooks/02_analises_descritivas.py"), encoding="utf-8").read()
_cols = _re.findall(r'"([A-Z_]+)"',
                    _re.search(r'COLUNAS_INFRA\s*(?::[^=]*)?=\s*\[(.*?)\]', _src, _re.S).group(1))
_pres = [c for c in _cols if c in painel.columns]
_p = painel.assign(indice_infra=painel[_pres].mean(axis=1))
_d24 = _p[_p["NU_ANO_CENSO"] == 2024]
checa("indice de infraestrutura 2024 — rural", 0.52,
      float(_d24[_d24["TP_LOCALIZACAO"] == 2]["indice_infra"].mean()), tol=0.02,
      fonte="painel + COLUNAS_INFRA do nb02")
checa("indice de infraestrutura 2024 — urbana", 0.71,
      float(_d24[_d24["TP_LOCALIZACAO"] == 1]["indice_infra"].mean()), tol=0.02,
      fonte="painel + COLUNAS_INFRA do nb02")

print()
print("=" * 78)
print("4. O ALVO (¶386, ¶387, ¶398)")
print("=" * 78)
alvo = feat["taxa_abandono_t1"]
checa("% observacoes com alvo zero", 68, (alvo == 0).mean() * 100, tol=0.02, fonte="features")
checa("% observacoes acima de 10% de abandono", 2, (alvo > 10).mean() * 100, tol=0.35,
      fonte="features")
checa("media do abandono anterior (Quadro 6)", 1.72, feat["abnd_t"].mean(), tol=0.01,
      fonte="features")
checa("maximo do abandono anterior", 43.6, feat["abnd_t"].max(), tol=0.01, fonte="features")

# grupos de localizacao em 2024 (taxa observada)
p24b = painel[painel["NU_ANO_CENSO"] == 2024]
jj = p24b.merge(t24, on="CO_ENTIDADE", how="inner")
if "TP_LOCALIZACAO_DIFERENCIADA" in jj.columns:
    dif = jj["TP_LOCALIZACAO_DIFERENCIADA"].fillna(0) > 0
    jj["grupo"] = np.where(dif, "diferenciada",
                           np.where(jj["TP_LOCALIZACAO"] == 2, "rural", "urbana"))
    g = jj.groupby("grupo")["TAXA_ABND_MED"].agg(["count", "mean", "std"])
    print("   grupos 2024 (recalculado):")
    print("   " + g.round(2).to_string().replace("\n", "\n   "))
    for nome, n_txt, m_txt in [("urbana", 679, 0.6), ("rural", 70, 0.9), ("diferenciada", 42, 6.3)]:
        if nome in g.index:
            checa(f"2024 {nome}: nº de escolas", n_txt, float(g.loc[nome, "count"]), tol=0.03,
                  fonte="painel x taxas")
            checa(f"2024 {nome}: abandono medio", m_txt, float(g.loc[nome, "mean"]), tol=0.09,
                  fonte="painel x taxas")

print()
print("=" * 78)
print("5. QUADROS 6, 7 e 8 (estatisticas das caracteristicas)")
print("=" * 78)
cont = est[est["tipo"] == "contínua"]
bina = est[est["tipo"].str.startswith("bin")]
checa("nº de caracteristicas continuas", 23, len(cont), fonte="estatisticas_features.csv")
checa("nº de caracteristicas binarias", 15, len(bina), fonte="estatisticas_features.csv")
sig = est[est["tipo"] != "categórica"]
checa("numericas com relacao significativa", 30, int((sig["p_valor"] < 0.05).sum()),
      fonte="estatisticas_features.csv")
checa("total de numericas", 38, len(sig), fonte="estatisticas_features.csv")

# confere linha a linha o Quadro 6 contra o CSV
q6 = TABS[6]
erros_q6 = []
for row in q6.rows[1:]:
    nome = row.cells[0].text.strip()
    linha = est[est["informacao"] == nome]
    if linha.empty:
        erros_q6.append((nome, "sem linha no CSV"))
        continue
    l = linha.iloc[0]
    for col, campo in [(1, "media"), (2, "desvio"), (3, "minimo"), (4, "mediana"),
                       (5, "maximo")]:
        try:
            v = float(row.cells[col].text.strip().replace(".", "").replace(",", "."))
        except ValueError:
            continue
        if abs(v - float(l[campo])) > max(0.011, abs(float(l[campo])) * 0.011):
            erros_q6.append((nome, f"{campo}: quadro={v} csv={float(l[campo]):.3f}"))
    rel = row.cells[6].text.strip().replace("−", "-").replace(",", ".").rstrip("*")
    try:
        v = float(rel)
        if abs(v - float(l["relacao"])) > 0.011:
            erros_q6.append((nome, f"relacao: quadro={v} csv={float(l['relacao']):.3f}"))
    except ValueError:
        pass
print(f"  Quadro 6 ({len(q6.rows)-1} linhas): {len(erros_q6)} divergencia(s)")
for e in erros_q6:
    print("    XX", e)
    res["div"].append(("Quadro 6 " + e[0], e[1], float("nan"), "estatisticas_features.csv"))
if not erros_q6:
    res["ok"] += len(q6.rows) - 1

q7 = TABS[7]
erros_q7 = []
for row in q7.rows[1:]:
    nome = row.cells[0].text.strip()
    linha = est[est["informacao"] == nome]
    if linha.empty:
        erros_q7.append((nome, "sem linha no CSV"))
        continue
    l = linha.iloc[0]
    def num(s):
        s = s.replace("−", "-").replace("%", "").replace(".", "").replace(",", ".")
        s = re.sub(r"\(.*?\)", "", s).strip().rstrip("*")
        return float(s)
    try:
        if abs(num(row.cells[1].text) - float(l["n_tem"])) > 0.5:
            erros_q7.append((nome, f"n_tem quadro={row.cells[1].text} csv={l['n_tem']}"))
        if abs(num(row.cells[3].text) - float(l["abandono_com"])) > 0.011:
            erros_q7.append((nome, f"abandono_com quadro={row.cells[3].text} csv={l['abandono_com']:.2f}"))
        if abs(num(row.cells[4].text) - float(l["abandono_sem"])) > 0.011:
            erros_q7.append((nome, f"abandono_sem quadro={row.cells[4].text} csv={l['abandono_sem']:.2f}"))
        if abs(num(row.cells[5].text) - float(l["diferenca"])) > 0.011:
            erros_q7.append((nome, f"diferenca quadro={row.cells[5].text} csv={l['diferenca']:.2f}"))
    except ValueError as e:
        erros_q7.append((nome, f"nao consegui ler: {e}"))
print(f"  Quadro 7 ({len(q7.rows)-1} linhas): {len(erros_q7)} divergencia(s)")
for e in erros_q7:
    print("    XX", e)
    res["div"].append(("Quadro 7 " + e[0], e[1], float("nan"), "estatisticas_features.csv"))
if not erros_q7:
    res["ok"] += len(q7.rows) - 1

q8 = TABS[8]
erros_q8 = []
for row in q8.rows[1:]:
    nome = row.cells[0].text.strip()
    l = meso[meso["categoria"] == nome]
    if l.empty:
        erros_q8.append((nome, "sem linha no CSV")); continue
    l = l.iloc[0]
    n = float(row.cells[1].text.replace(".", ""))
    pct = float(row.cells[2].text.replace("%", "").replace(",", "."))
    ab = float(row.cells[3].text.replace("%", "").replace(",", "."))
    if abs(n - l["escolas"]) > 0.5:
        erros_q8.append((nome, f"n quadro={n} csv={l['escolas']}"))
    if abs(pct - l["participacao"]) > 0.06:
        erros_q8.append((nome, f"% quadro={pct} csv={l['participacao']:.2f}"))
    if abs(ab - l["abandono_medio"]) > 0.011:
        erros_q8.append((nome, f"abandono quadro={ab} csv={l['abandono_medio']:.2f}"))
print(f"  Quadro 8 ({len(q8.rows)-1} linhas): {len(erros_q8)} divergencia(s)")
for e in erros_q8:
    print("    XX", e)
    res["div"].append(("Quadro 8 " + e[0], e[1], float("nan"), "estatisticas_mesorregiao.csv"))
if not erros_q8:
    res["ok"] += len(q8.rows) - 1

print()
print("=" * 78)
print("6. QUADRO 9 (validacao repetida) e QUADRO 10 (teste temporal)")
print("=" * 78)
MAP9 = {"Sem sistema (média da rede)": "dummy", "Modelo de referência linear": "ridge",
        "Modelo de referência em árvores": "random_forest", "Modelo adotado": "xgboost"}
q9 = TABS[9]
for row in q9.rows[1:]:
    nome = row.cells[0].text.strip()
    mod = MAP9.get(nome)
    if not mod:
        continue
    d = cvrep[cvrep["modelo"] == mod]
    for col, campo in [(1, "rmse"), (2, "spearman"), (3, "precision_at_k")]:
        celula = row.cells[col].text.strip()
        m = re.match(r"^([\d,]+)", celula)
        if not m:
            continue
        v = float(m.group(1).replace(",", "."))
        calc = d[campo].mean()
        if pd.isna(calc):
            continue
        ok = abs(v - calc) <= max(0.0011, abs(calc) * 0.012)
        print(f"  {'ok' if ok else 'XX'} Q9 {nome} / {campo}: quadro={v} calc={calc:.4f}")
        (res.__setitem__("ok", res["ok"] + 1) if ok
         else res["div"].append((f"Q9 {nome}/{campo}", v, calc, "metricas_cv_repetida.csv")))

MAP10 = dict(MAP9); MAP10["Sem sistema (média da rede)"] = "dummy_media"
q10 = TABS[10]
mbas = pd.read_csv(os.path.join(RAIZ, "reports/metricas_baselines.csv"))
tmp = pd.concat([mxgb, mbas])
tmp = tmp[tmp["avaliacao"] == "temporal_2023_2024"]
for row in q10.rows[1:]:
    nome = row.cells[0].text.strip()
    mod = MAP10.get(nome)
    if not mod:
        continue
    d = tmp[tmp["modelo"] == mod]
    if d.empty:
        print(f"  ?? Q10 {nome}: modelo '{mod}' ausente do CSV temporal")
        res["na"].append((f"Q10 {nome}", "-", "metricas_xgboost.csv"))
        continue
    for col, campo in [(1, "rmse"), (2, "spearman"), (3, "precision_at_k"), (4, "roc_auc")]:
        if campo not in d.columns:
            continue
        celula = row.cells[col].text.strip()
        m = re.match(r"^([\d,]+)", celula)
        if not m:
            continue
        v = float(m.group(1).replace(",", "."))
        calc = float(d[campo].mean())
        if pd.isna(calc):
            continue
        ok = abs(v - calc) <= max(0.0011, abs(calc) * 0.012)
        print(f"  {'ok' if ok else 'XX'} Q10 {nome} / {campo}: quadro={v} calc={calc:.4f}")
        (res.__setitem__("ok", res["ok"] + 1) if ok
         else res["div"].append((f"Q10 {nome}/{campo}", v, calc, "metricas_xgboost.csv")))

print()
print("=" * 78)
print("7. QUADRO 11 e transicoes anuais (SS13)")
print("=" * 78)
g = ss13.groupby("cenario")[["rmse", "spearman", "precision_at_k", "roc_auc"]].mean()
print("   recalculado:")
print("   " + g.round(3).to_string().replace("\n", "\n   "))
q11 = TABS[11]
CAMPOS = {"Erro médio": "rmse", "Acerto da ordenação": "spearman",
          "Acerto na lista": "precision_at_k", "Separação alto": "roc_auc"}
cen = list(g.index)
for row in q11.rows[1:]:
    rot = row.cells[0].text.strip()
    campo = next((v for k, v in CAMPOS.items() if rot.startswith(k)), None)
    if not campo:
        continue
    for i, c in enumerate(cen, start=1):
        m = re.match(r"^([\d,]+)", row.cells[i].text.strip())
        if not m:
            continue
        v = float(m.group(1).replace(",", "."))
        calc = float(g.loc[c, campo])
        ok = abs(v - calc) <= max(0.006, abs(calc) * 0.012)
        print(f"  {'ok' if ok else 'XX'} Q11 {campo} / {c}: quadro={v} calc={calc:.4f}")
        (res.__setitem__("ok", res["ok"] + 1) if ok
         else res["div"].append((f"Q11 {campo}/{c}", v, calc, "metricas_ss13_cenarios.csv")))

# ¶452: 63% e 38% de acerto na lista; abandono medio 1,36% -> 0,88%
checa("¶452 acerto da lista 2022→2023", 63, float(g.loc[cen[0], "precision_at_k"]) * 100,
      tol=0.02, fonte="ss13")
checa("¶452 acerto da lista 2023→2024", 38, float(g.loc[cen[1], "precision_at_k"]) * 100,
      tol=0.03, fonte="ss13")
checa("¶452 ROC 2023→2024", 0.81, float(g.loc[cen[1], "roc_auc"]), tol=0.02, fonte="ss13")
m22 = feat[feat["NU_ANO_CENSO"] == 2022]["taxa_abandono_t1"].mean()
m23 = feat[feat["NU_ANO_CENSO"] == 2023]["taxa_abandono_t1"].mean()
checa("¶452 abandono medio da coorte 2023", 1.36, m22, tol=0.02, fonte="features")
checa("¶452 abandono medio da coorte 2024", 0.88, m23, tol=0.02, fonte="features")

print()
print("=" * 78)
print("8. RESIDUOS EM ALUNOS, EQUIDADE e DIAGNOSTICO")
print("=" * 78)
mae = ralunos.groupby("transicao")["residuo_alunos"].apply(lambda s: s.abs().mean())
print("   MAE por transicao:", {k: round(v, 2) for k, v in mae.items()})
vals = sorted(mae.values, reverse=True)
checa("MAE em alunos — transicao 1", 3.3, vals[0], tol=0.03, fonte="residuos_transicoes_alunos.csv")
checa("MAE em alunos — transicao 2", 3.0, vals[1], tol=0.03, fonte="residuos_transicoes_alunos.csv")

eq = rgrupo.groupby("grupo")["residuo"].mean()
checa("equidade diferenciada", -6.12, float(eq["diferenciada"]), tol=0.01, fonte="residuos_grupo_temporal.csv")
checa("equidade urbana", 0.07, float(eq["urbana"]), tol=0.02, fonte="residuos_grupo_temporal.csv")
checa("equidade rural", 0.38, float(eq["rural"]), tol=0.02, fonte="residuos_grupo_temporal.csv")

col_r = "residuo" if "residuo" in rdiag.columns else rdiag.columns[-1]
checa("¶505 desvio minimo (P3)", -27.87, float(rdiag[col_r].min()), tol=0.01,
      fonte="residuos_top20_temporal.csv")
checa("¶505 desvio maximo (P3)", 23.69, float(rdiag[col_r].max()), tol=0.01,
      fonte="residuos_top20_temporal.csv")

print()
print("=" * 78)
print("9. SHAP: ordem dos fatores citada no ¶471")
print("=" * 78)
top6 = list(shap_imp.head(6)["feature"])
print("   top-6 no CSV:", top6)
texto467 = next(p for p in PARS if "o melhor previsor do tempo de amanhã" in p)
print("   texto: ...", texto467[texto467.find("ela mostra"):][:230])

print()
print("=" * 78)
print("10. SUITE DE TESTES")
print("=" * 78)
print("   (contagem verificada por execucao do pytest, em separado)")

print()
print("=" * 78)
print(f"RESUMO: {res['ok']} conferidos OK · {len(res['div'])} divergencia(s) · "
      f"{len(res['na'])} sem fonte automatica")
print("=" * 78)
for d in res["div"]:
    print("  DIVERGE:", d)
for n in res["na"]:
    print("  SEM FONTE:", n)
